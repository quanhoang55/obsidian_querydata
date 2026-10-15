"""FastAPI endpoints for local folder selection and Gemini queries."""

from __future__ import annotations

import asyncio
import logging
from dataclasses import replace
from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import Any, Literal

from fastapi import APIRouter, HTTPException, Path as PathParam
from fastapi.encoders import jsonable_encoder
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, SecretStr, field_validator

from backend.clients.gemini import GeminiClientError, GeminiTimeoutError
from backend.config import Settings, get_settings, set_runtime_settings
from backend.schemas.query import DisplayPlan, MultiFileQueryPlan
from backend.services.app_settings_store import (
    SettingsStorageError,
    read_preferences,
    set_api_key,
    write_preferences,
)
from backend.services.local_vault import LocalVault, VaultNotSelectedError, select_folder, validate_root
from backend.services.query_compiler import InvalidQueryPlan
from backend.services.query_executor import DuckDBUnavailableError, QueryExecutionError
from backend.services.query_service import VaultConfigurationError, query_files


logger = logging.getLogger(__name__)


class HealthResponse(BaseModel):
    status: Literal["ok"]


class ConfigResponse(BaseModel):
    vault_path: str | None
    gemini_api_key_configured: bool
    port: int


class ConfigUpdate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    vault_path: str = Field(min_length=1)

    @field_validator("vault_path")
    @classmethod
    def validate_vault_path(cls, value: str) -> str:
        try:
            return str(validate_root(value.strip()))
        except (OSError, RuntimeError, ValueError) as error:
            raise ValueError("vault_path must point to an existing readable folder") from error


class SettingsResponse(BaseModel):
    gemini_api_key_configured: bool
    gemini_model: str
    gemini_timeout: float


class SettingsUpdate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    gemini_api_key: SecretStr | None = None
    gemini_model: str | None = Field(default=None, min_length=1, max_length=200)
    gemini_timeout: float | None = Field(default=None, gt=0, le=600)
    clear_gemini_api_key: bool = False

    @field_validator("gemini_api_key")
    @classmethod
    def validate_api_key(cls, value: SecretStr | None) -> SecretStr | None:
        if value is None:
            return None
        secret = value.get_secret_value().strip()
        if not secret or len(secret) > 512:
            raise ValueError("Gemini API key must contain 1 to 512 characters")
        return SecretStr(secret)

    @field_validator("gemini_model")
    @classmethod
    def trim_model(cls, value: str | None) -> str | None:
        if value is None:
            return None
        value = value.strip()
        if not value:
            raise ValueError("gemini_model must not be blank")
        return value


class VaultFilesResponse(BaseModel):
    files: list[str]
    total: int


class SearchRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    query: str = Field(min_length=1, max_length=500)
    context_length: int = Field(default=100, ge=0, le=1000)

    @field_validator("query")
    @classmethod
    def trim_query(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("query must not be blank")
        return value


class QueryRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    prompt: str = Field(min_length=1, max_length=2000)

    @field_validator("prompt")
    @classmethod
    def trim_prompt(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("prompt must not be blank")
        return value


class ResultColumn(BaseModel):
    name: str
    type: str


class QueryResultData(BaseModel):
    columns: list[ResultColumn]
    rows: list[dict[str, Any]]
    row_count: int
    truncated: bool


class QueryResponse(BaseModel):
    status: Literal["success", "no_results", "clarification_required"]
    query: str
    answer: str
    plan: MultiFileQueryPlan
    display: DisplayPlan
    result: QueryResultData
    step_results: list[dict[str, Any]]
    sources: list[str]
    warnings: list[dict[str, str]]


router = APIRouter()


def _public_config(settings: Settings) -> ConfigResponse:
    return ConfigResponse(
        vault_path=settings.vault_path,
        gemini_api_key_configured=bool(settings.gemini_api_key),
        port=settings.port,
    )


def _local_vault() -> LocalVault:
    try:
        return LocalVault(get_settings().vault_path)
    except VaultNotSelectedError as error:
        raise HTTPException(status_code=409, detail=str(error)) from error


def _excel_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time, timedelta)):
        return jsonable_encoder(value)
    return value


@router.get("/health", response_model=HealthResponse, tags=["system"])
async def health() -> HealthResponse:
    return HealthResponse(status="ok")


@router.get("/api/v1/config", response_model=ConfigResponse, tags=["config"])
async def read_config() -> ConfigResponse:
    return _public_config(get_settings())


@router.put("/api/v1/config", response_model=ConfigResponse, tags=["config"])
async def update_config(payload: ConfigUpdate) -> ConfigResponse:
    """Set a selected directory as the root for all local file operations."""
    try:
        preferences = read_preferences()
        preferences["vault_path"] = payload.vault_path
        write_preferences(preferences)
    except SettingsStorageError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error
    updated = replace(get_settings(), vault_path=payload.vault_path)
    set_runtime_settings(updated)
    return _public_config(updated)


@router.get("/api/v1/settings", response_model=SettingsResponse, tags=["settings"])
async def read_settings() -> SettingsResponse:
    try:
        settings = get_settings()
    except SettingsStorageError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error
    return SettingsResponse(
        gemini_api_key_configured=bool(settings.gemini_api_key),
        gemini_model=settings.gemini_model,
        gemini_timeout=settings.gemini_timeout,
    )


@router.put("/api/v1/settings", response_model=SettingsResponse, tags=["settings"])
async def update_settings(payload: SettingsUpdate) -> SettingsResponse:
    """Persist Gemini preferences and keep the API key in the OS credential store."""
    if payload.gemini_api_key is not None and payload.clear_gemini_api_key:
        raise HTTPException(status_code=422, detail="Set a Gemini API key or clear it, not both")
    try:
        preferences = read_preferences()
        current = get_settings()
        next_key = current.gemini_api_key
        if payload.gemini_api_key is not None:
            next_key = payload.gemini_api_key.get_secret_value()
            set_api_key(next_key)
            preferences["gemini_api_key_cleared"] = False
        elif payload.clear_gemini_api_key:
            set_api_key(None)
            next_key = ""
            preferences["gemini_api_key_cleared"] = True

        next_model = payload.gemini_model or current.gemini_model
        next_timeout = payload.gemini_timeout or current.gemini_timeout
        if payload.gemini_model is not None:
            preferences["gemini_model"] = next_model
        if payload.gemini_timeout is not None:
            preferences["gemini_timeout"] = next_timeout
        write_preferences(preferences)
    except SettingsStorageError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error

    updated = replace(
        current,
        gemini_api_key=next_key,
        gemini_model=next_model,
        gemini_timeout=next_timeout,
    )
    set_runtime_settings(updated)
    return SettingsResponse(
        gemini_api_key_configured=bool(updated.gemini_api_key),
        gemini_model=updated.gemini_model,
        gemini_timeout=updated.gemini_timeout,
    )


@router.post("/api/v1/vault/select", tags=["vault"])
async def choose_vault_folder() -> dict[str, Any]:
    """Open a native folder chooser on the computer running this backend."""
    try:
        selected = await asyncio.to_thread(select_folder)
    except RuntimeError as error:
        raise HTTPException(status_code=500, detail=str(error)) from error
    if selected is None:
        return {"cancelled": True, "vault_path": get_settings().vault_path}
    try:
        preferences = read_preferences()
        preferences["vault_path"] = selected
        write_preferences(preferences)
    except SettingsStorageError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error
    updated = replace(get_settings(), vault_path=selected)
    set_runtime_settings(updated)
    return {"cancelled": False, **_public_config(updated).model_dump()}


@router.get("/api/v1/vault/files", response_model=VaultFilesResponse, tags=["vault"])
async def list_vault_files() -> VaultFilesResponse:
    try:
        files = await asyncio.to_thread(_local_vault().list_files)
    except HTTPException:
        raise
    except OSError as error:
        raise HTTPException(status_code=500, detail="Could not list files in the selected folder") from error
    return VaultFilesResponse(files=files, total=len(files))


@router.get("/api/v1/vault/files/{path:path}", tags=["vault"])
async def read_vault_file(
    path: str = PathParam(..., description="Path relative to the selected folder"),
) -> dict[str, Any]:
    try:
        vault = _local_vault()
        file_path = vault.resolve_file(path)
        content = await asyncio.to_thread(file_path.read_bytes)
        relative_path = vault.relative(file_path)
    except HTTPException:
        raise
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    except FileNotFoundError as error:
        raise HTTPException(status_code=404, detail="Supported file not found in selected folder") from error
    if relative_path.lower().endswith(".md"):
        try:
            return {"path": relative_path, "type": "markdown", "content": content.decode("utf-8")}
        except UnicodeDecodeError as error:
            raise HTTPException(status_code=422, detail="Markdown file is not valid UTF-8") from error
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            sheets = [
                {
                    "name": worksheet.title,
                    "rows": [
                        [_excel_value(value) for value in row]
                        for row in worksheet.iter_rows(values_only=True)
                    ],
                }
                for worksheet in workbook.worksheets
            ]
        finally:
            workbook.close()
    except Exception as error:
        raise HTTPException(status_code=422, detail="Excel file could not be parsed") from error
    return {"path": relative_path, "type": "excel", "sheets": sheets}


@router.post("/api/v1/vault/search", tags=["vault"])
async def search_vault(payload: SearchRequest) -> dict[str, Any]:
    try:
        results = await asyncio.to_thread(_local_vault().search, payload.query, payload.context_length)
    except HTTPException:
        raise
    except UnicodeDecodeError as error:
        raise HTTPException(status_code=422, detail="A Markdown file is not valid UTF-8") from error
    except (OSError, ValueError) as error:
        raise HTTPException(status_code=500, detail="Could not search the selected folder") from error
    return {"query": payload.query, "results": results, "total": len(results)}


@router.post("/api/v1/query", response_model=QueryResponse, tags=["query"])
async def query_vault(payload: QueryRequest) -> QueryResponse:
    settings = get_settings()
    if not settings.gemini_api_key:
        raise HTTPException(status_code=503, detail="Set GEMINI_API_KEY to enable natural language queries")
    try:
        return QueryResponse.model_validate(await query_files(payload.prompt, settings))
    except VaultConfigurationError as error:
        raise HTTPException(status_code=409, detail=str(error)) from error
    except GeminiTimeoutError as error:
        raise HTTPException(status_code=504, detail="Gemini query interpretation timed out") from error
    except GeminiClientError as error:
        raise HTTPException(status_code=502, detail=str(error)) from error
    except InvalidQueryPlan as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    except DuckDBUnavailableError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error
    except QueryExecutionError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    except UnicodeDecodeError as error:
        raise HTTPException(status_code=422, detail="A Markdown file is not valid UTF-8") from error
    except (OSError, ValueError) as error:
        raise HTTPException(status_code=422, detail="Could not parse local vault data") from error
    except Exception as error:
        # Always return JSON to the frontend. The traceback remains in the local
        # Uvicorn console for diagnosis without exposing file-system details.
        logger.exception("Unexpected error while processing a vault query")
        raise HTTPException(status_code=500, detail="Unexpected error while processing the selected folder") from error


@router.get("/api/v1/query/examples", tags=["query"])
async def query_examples() -> dict[str, list[str]]:
    return {
        "examples": [
            "Tổng doanh số tháng này",
            "Tổng doanh số của Nam từ ngày 01/05/2026 đến 10/05/2026",
            "Doanh số của Lan trong tuần trước",
        ]
    }


@router.get("/api/v1/tags", tags=["vault"])
async def list_vault_tags() -> dict[str, Any]:
    try:
        tags = await asyncio.to_thread(_local_vault().tags)
    except HTTPException:
        raise
    except UnicodeDecodeError as error:
        raise HTTPException(status_code=422, detail="A Markdown file is not valid UTF-8") from error
    except (OSError, ValueError) as error:
        raise HTTPException(status_code=500, detail="Could not read tags from the selected folder") from error
    return {"tags": tags, "total": len(tags)}
