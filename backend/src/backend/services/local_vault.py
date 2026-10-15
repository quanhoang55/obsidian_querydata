"""Read supported files under the user-selected local workspace folder."""

from __future__ import annotations

import os
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SUPPORTED_SUFFIXES = {".md", ".xlsx"}
_FOLDER_DIALOG = """
import tkinter as tk
from tkinter import filedialog
root = tk.Tk()
root.withdraw()
root.attributes('-topmost', True)
path = filedialog.askdirectory(title='Choose an Obsidian vault or data folder')
root.destroy()
print(path)
"""


class VaultNotSelectedError(Exception):
    """No local workspace folder has been selected."""


def select_folder() -> str | None:
    """Open the operating system's native folder chooser."""
    try:
        result = subprocess.run(
            [sys.executable, "-c", _FOLDER_DIALOG],
            check=False,
            capture_output=True,
            text=True,
            timeout=180,
        )
    except (OSError, subprocess.TimeoutExpired) as error:
        raise RuntimeError("Could not open the native folder picker") from error
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or "Folder picker failed")
    selected = result.stdout.strip()
    return str(validate_root(selected)) if selected else None


def validate_root(value: str | os.PathLike[str]) -> Path:
    """Resolve and verify a readable directory selected by the user."""
    root = Path(value).expanduser().resolve(strict=True)
    if not root.is_dir():
        raise ValueError("Selected path is not a folder")
    if not os.access(root, os.R_OK | os.X_OK):
        raise ValueError("Selected folder is not readable")
    return root


class LocalVault:
    """Provide path-safe listing, file reads, text search, and tags."""

    def __init__(self, root_path: str | None) -> None:
        if not root_path:
            raise VaultNotSelectedError("Choose a vault folder before using this feature")
        try:
            self.root = validate_root(root_path)
        except (OSError, RuntimeError, ValueError) as error:
            raise VaultNotSelectedError("The selected vault folder is unavailable") from error

    def resolve_file(self, relative_path: str) -> Path:
        normalized = relative_path.replace("\\", "/")
        parts = normalized.split("/")
        if not normalized or normalized.startswith("/") or any(part in {"", ".", ".."} for part in parts):
            raise ValueError("Path must stay inside the selected folder")
        candidate = (self.root / Path(*parts)).resolve(strict=True)
        if not candidate.is_relative_to(self.root):
            raise ValueError("Path must stay inside the selected folder")
        if not candidate.is_file() or candidate.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise FileNotFoundError(relative_path)
        return candidate

    def relative(self, path: Path) -> str:
        return path.relative_to(self.root).as_posix()

    def list_files(self) -> list[str]:
        files: list[str] = []
        for directory, directories, filenames in os.walk(self.root, followlinks=False):
            parent = Path(directory)
            directories[:] = [
                name for name in directories
                if not name.startswith(".") and not (parent / name).is_symlink()
            ]
            for filename in filenames:
                path = parent / filename
                if filename.startswith(".") or path.is_symlink() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
                    continue
                files.append(self.relative(path))
        return sorted(files, key=str.casefold)

    def read_bytes(self, relative_path: str) -> bytes:
        return self.resolve_file(relative_path).read_bytes()

    def documents(self) -> list[tuple[str, str | bytes]]:
        documents: list[tuple[str, str | bytes]] = []
        for relative_path in self.list_files():
            path = self.resolve_file(relative_path)
            content = path.read_bytes()
            documents.append((relative_path, content.decode("utf-8") if path.suffix.lower() == ".md" else content))
        return documents

    def search(self, query: str, context_length: int) -> list[dict[str, Any]]:
        needle = query.casefold()
        results: list[dict[str, Any]] = []
        for relative_path in self.list_files():
            path = self.resolve_file(relative_path)
            if path.suffix.lower() == ".md":
                lines = path.read_text(encoding="utf-8").splitlines()
                for line_number, line in enumerate(lines, start=1):
                    if needle in line.casefold():
                        results.append({
                            "file": relative_path,
                            "line": line_number,
                            "context": line[:context_length] if context_length else "",
                        })
            else:
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    for worksheet in workbook.worksheets:
                        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                            line = " | ".join(str(value) for value in row if value is not None)
                            if needle in line.casefold():
                                results.append({
                                    "file": relative_path,
                                    "sheet": worksheet.title,
                                    "row": row_number,
                                    "context": line[:context_length] if context_length else "",
                                })
                finally:
                    workbook.close()
        return results

    def tags(self) -> list[dict[str, Any]]:
        counts: Counter[str] = Counter()
        pattern = re.compile(r"(?<![\w/])#([\w/-]+)")
        for relative_path in self.list_files():
            path = self.resolve_file(relative_path)
            if path.suffix.lower() == ".md":
                counts.update(match.group(1) for match in pattern.finditer(path.read_text(encoding="utf-8")))
        return [{"name": name, "count": count} for name, count in sorted(counts.items(), key=lambda item: item[0].casefold())]
