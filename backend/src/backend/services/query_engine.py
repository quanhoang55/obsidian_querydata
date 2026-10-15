"""Parse local Markdown/Excel tables and execute a validated query intent."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


HEADER_GROUPS = {
    "date": ("date", "ngay", "ngay ban", "ngay tao", "created", "day", "thoi gian"),
    "person": (
        "salesperson", "sales person", "seller", "seller name", "sale", "sales", "sale person", "sale name", "employee", "employee name",
        "staff", "staff name", "sales staff", "sales employee", "sales rep", "representative",
        "account manager", "consultant", "name", "nv", "nvkd", "nvbh", "nhanvien", "nhan vien", "nhan vien sale",
        "nhan vien ban hang", "nguoi ban", "nguoi phu trach", "ten", "ten nhan vien", "ten nv", "ten nvkd", "ten nvbh", "ten sale",
    ),
    "product": (
        "product", "product name", "item", "item name", "mat hang", "san pham", "ten mat hang", "ten san pham",
        "ten hang", "hang hoa", "sku name",
    ),
    "amount": (
        "doanh so", "doanh thu", "revenue", "sales amount", "amount", "thanh tien",
        "tong tien", "gia tri", "total", "value", "so tien", "thanh toan",
    ),
}
DATE_PATTERN = re.compile(r"(?<!\d)(\d{1,2}[/-]\d{1,2}[/-]\d{4}|\d{4}-\d{2}-\d{2})(?!\d)")


@dataclass(frozen=True, slots=True)
class TableRow:
    path: str
    values: dict[str, Any]


def normalize_text(value: str) -> str:
    value = value.casefold().replace("đ", "d")
    decomposed = unicodedata.normalize("NFKD", value)
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def _clean_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values):
        label = str(value).strip() if value is not None else ""
        label = label or f"Column {index + 1}"
        seen[label] = seen.get(label, 0) + 1
        headers.append(f"{label} ({seen[label]})" if seen[label] > 1 else label)
    return headers


def _header_group(label: str) -> str | None:
    normalized = normalize_text(label).strip()
    exact = [group for group, aliases in HEADER_GROUPS.items() if normalized in aliases]
    if exact:
        return exact[0]
    # Prefer the most specific prefix. For example, "Tên mặt hàng 1" must
    # match "tên mặt hàng" (product), rather than the shorter "tên" (person).
    matches = [
        (len(alias), group)
        for group, aliases in HEADER_GROUPS.items()
        for alias in aliases
        if normalized.startswith(f"{alias} ")
    ]
    return max(matches)[1] if matches else None


def _resolve_column(headers: list[str], requested: str | None, group: str | None = None) -> str | None:
    if requested:
        normalized = normalize_text(requested).strip()
        for header in headers:
            if normalize_text(header).strip() == normalized:
                return header
        requested_group = _header_group(requested)
        if requested_group:
            group = requested_group
    if group:
        for header in headers:
            if _header_group(header) == group:
                return header
    return None


def _rows_from_values(path: str, headers: list[str], source_rows: list[tuple[Any, ...]]) -> list[TableRow]:
    rows: list[TableRow] = []
    for source in source_rows:
        if not any(value is not None and str(value).strip() for value in source):
            continue
        values = {header: source[index] if index < len(source) else None for index, header in enumerate(headers)}
        rows.append(TableRow(path, values))
    return rows


def parse_excel_records(path: str, content: bytes) -> list[TableRow]:
    """Read all non-empty worksheet rows as column/value maps."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    records: list[TableRow] = []
    try:
        for worksheet in workbook.worksheets:
            raw_rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
            nonempty_indexes = [
                index for index, row in enumerate(raw_rows)
                if any(value is not None and str(value).strip() for value in row)
            ]
            if not nonempty_indexes:
                continue
            first_index = nonempty_indexes[0]
            first_values = [value for value in raw_rows[first_index] if value is not None and str(value).strip()]
            first_is_header = len(first_values) >= 2 or any(_header_group(str(value)) for value in first_values)
            header_index = first_index if first_is_header else next(
                (
                    index for index in nonempty_indexes[1:]
                    if any(_header_group(str(value)) for value in raw_rows[index] if value is not None)
                    or sum(value is not None and bool(str(value).strip()) for value in raw_rows[index]) >= 2
                ),
                first_index,
            )
            headers = _clean_headers(list(raw_rows[header_index]))
            records.extend(_rows_from_values(path, headers, raw_rows[header_index + 1 :]))
    finally:
        workbook.close()
    return records


def describe_documents(documents: list[tuple[str, str | bytes]]) -> list[dict[str, Any]]:
    """Return table headers only, so Gemini can ground intent in real schemas."""
    descriptions: list[dict[str, Any]] = []
    for path, content in documents:
        if path.lower().endswith(".xlsx") and isinstance(content, bytes):
            sheets: list[dict[str, Any]] = []
            try:
                workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
                try:
                    for worksheet in workbook.worksheets:
                        # Header discovery only needs the opening rows; avoid materializing
                        # thousands of transaction rows just to build Gemini's schema context.
                        raw_rows = [
                            tuple(row)
                            for row in worksheet.iter_rows(min_row=1, max_row=30, values_only=True)
                        ]
                        nonempty_indexes = [
                            index for index, row in enumerate(raw_rows)
                            if any(value is not None and str(value).strip() for value in row)
                        ]
                        if not nonempty_indexes:
                            continue
                        first_index = nonempty_indexes[0]
                        first_values = [value for value in raw_rows[first_index] if value is not None and str(value).strip()]
                        first_is_header = len(first_values) >= 2 or any(_header_group(str(value)) for value in first_values)
                        header_index = first_index if first_is_header else next(
                            (
                                index for index in nonempty_indexes[1:]
                                if any(_header_group(str(value)) for value in raw_rows[index] if value is not None)
                                or sum(value is not None and bool(str(value).strip()) for value in raw_rows[index]) >= 2
                            ),
                            first_index,
                        )
                        sheets.append({"name": worksheet.title, "columns": _clean_headers(list(raw_rows[header_index]))})
                finally:
                    workbook.close()
            except Exception:
                # The execution phase will report this workbook in `warnings`.
                # A damaged or unsupported workbook must not make every query fail.
                continue
            descriptions.append({"file": path, "sheets": sheets})
        elif path.lower().endswith(".md") and isinstance(content, str):
            rows = parse_markdown_records(path, content)
            columns = list(dict.fromkeys(column for row in rows for column in row.values))
            if columns:
                descriptions.append({"file": path, "sheets": [{"name": "Markdown tables", "columns": columns}]})
    return descriptions


def parse_markdown_records(path: str, content: str) -> list[TableRow]:
    """Read Markdown tables into the same generic row representation."""
    lines = [line.strip() for line in content.splitlines() if line.strip()]
    records: list[TableRow] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if "|" not in line:
            index += 1
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        separator_follows = index + 1 < len(lines) and re.fullmatch(r"[|\s:-]+", lines[index + 1])
        recognized_header = any(_header_group(cell) for cell in cells)
        if not separator_follows and not recognized_header:
            index += 1
            continue
        headers = _clean_headers(cells)
        index += 2 if separator_follows else 1
        while index < len(lines) and "|" in lines[index]:
            values = tuple(cell.strip() for cell in lines[index].strip("|").split("|"))
            if not re.fullmatch(r"[|\s:-]+", lines[index]):
                records.extend(_rows_from_values(path, headers, [values]))
            index += 1
    return records


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        match = DATE_PATTERN.search(value.strip())
        if not match:
            return None
        token = match.group(1)
        try:
            if re.match(r"\d{4}-", token):
                return date.fromisoformat(token)
            day, month, year = map(int, re.split(r"[/-]", token))
            return date(year, month, day)
        except ValueError:
            return None
    return None


def _as_number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().lower().replace("₫", "").replace("vnd", "").replace(",", "")
    multiplier = 1.0
    for suffix, amount in (("tỷ", 1e9), ("triệu", 1e6), ("million", 1e6), ("billion", 1e9), ("k", 1e3)):
        if suffix in raw:
            raw = raw.replace(suffix, "")
            multiplier = amount
            break
    raw = re.sub(r"[^0-9.+-]", "", raw)
    try:
        return float(raw) * multiplier
    except ValueError:
        return None


def _matches(row: TableRow, person: str | None, start_date: date | None, end_date: date | None) -> bool:
    headers = list(row.values)
    if person:
        person_column = _resolve_column(headers, None, "person")
        value = str(row.values.get(person_column, "") if person_column else "")
        if normalize_text(person) not in normalize_text(value):
            return False
    if start_date or end_date:
        date_column = _resolve_column(headers, None, "date")
        row_date = _as_date(row.values.get(date_column)) if date_column else None
        if row_date is None or (start_date and row_date < start_date) or (end_date and row_date > end_date):
            return False
    return True


def aggregate_sales(prompt: str, documents: list[tuple[str, str | bytes]], plan: dict[str, Any] | None = None) -> dict[str, Any]:
    """Execute aggregate/distinct/inspect intents against local table rows."""
    query_plan = plan or {}
    start_date = date.fromisoformat(query_plan["start_date"]) if query_plan.get("start_date") else None
    end_date = date.fromisoformat(query_plan["end_date"]) if query_plan.get("end_date") else None
    person = query_plan.get("person")
    operation = query_plan.get("operation", "aggregate")
    metric = query_plan.get("metric") or "sum"
    requested_path = query_plan.get("file_path")
    if not requested_path:
        path_match = re.search(r"((?:[\w().-]+/)*[\w().-]+\.(?:xlsx|md))\b", prompt, re.IGNORECASE)
        requested_path = path_match.group(1) if path_match else None
    if isinstance(requested_path, str):
        requested_path = requested_path.replace("\\", "/").casefold()

    records: list[TableRow] = []
    warnings: list[dict[str, str]] = []
    files_scanned = 0
    for path, content in documents:
        if requested_path and path.replace("\\", "/").casefold() != requested_path:
            continue
        if path.lower().endswith(".xlsx") and isinstance(content, bytes):
            files_scanned += 1
            try:
                records.extend(parse_excel_records(path, content))
            except Exception as error:
                warnings.append({"file": path, "message": f"Spreadsheet could not be parsed ({type(error).__name__})."})
        elif path.lower().endswith(".md") and isinstance(content, str):
            files_scanned += 1
            records.extend(parse_markdown_records(path, content))

    matched = [row for row in records if _matches(row, person, start_date, end_date)]
    all_headers = list(dict.fromkeys(header for row in matched for header in row.values))
    target_column = _resolve_column(all_headers, query_plan.get("target_column"))
    if target_column is None and operation == "aggregate":
        target_column = _resolve_column(all_headers, None, "amount")
    unresolved_target = bool(query_plan.get("target_column")) and target_column is None
    if unresolved_target:
        requested_name = str(query_plan["target_column"])
        warnings.append({
            "file": requested_path or (matched[0].path if matched else "vault"),
            "message": f"Could not match column '{requested_name}'. Available columns: {', '.join(all_headers) or 'none'}.",
        })

    value = 0.0
    distinct_values: list[Any] = []
    if operation == "distinct" and target_column:
        for row in matched:
            item = row.values.get(target_column)
            if item is not None and str(item).strip() and item not in distinct_values:
                distinct_values.append(item)
    elif operation == "aggregate":
        if metric == "count":
            value = float(sum(row.values.get(target_column) is not None for row in matched)) if target_column else float(len(matched))
        else:
            numbers = [_as_number(row.values.get(target_column)) for row in matched] if target_column else []
            numeric_values = [number for number in numbers if number is not None]
            value = sum(numeric_values)
            if metric == "average" and numeric_values:
                value /= len(numeric_values)

    display_columns: list[str] = []
    for requested in query_plan.get("display_columns", []):
        column = _resolve_column(all_headers, requested)
        if column and column not in display_columns:
            display_columns.append(column)
    if operation == "distinct" and target_column:
        display_columns = [target_column]
    elif not display_columns and target_column:
        display_columns = [target_column]

    if operation == "distinct":
        display_file = requested_path or (matched[0].path if matched else "")
        display_rows = [{"file": display_file, "values": {target_column: item}} for item in distinct_values] if target_column else []
        target_name = target_column or query_plan.get("target_column") or "requested column"
        if unresolved_target:
            answer = f"Không tìm thấy cột '{target_name}' trong dữ liệu. Các cột có sẵn: {', '.join(all_headers) or 'không có cột nào'}."
        else:
            answer = f"Các giá trị khác nhau trong cột {target_name}: " + (", ".join(map(str, distinct_values)) if distinct_values else "không tìm thấy dữ liệu phù hợp.")
        total: float | int = len(distinct_values)
    else:
        display_rows = [
            {"file": row.path, "values": {column: row.values.get(column) for column in display_columns}}
            for row in matched
        ]
        total = value
        grouped_results: dict[str, float] = {}
        grouping = query_plan.get("group_by", "none")
        if operation == "aggregate" and grouping != "none":
            if grouping == "person":
                group_column = _resolve_column(all_headers, None, "person")
                keys = [str(row.values.get(group_column) or "Unknown") for row in matched]
            elif grouping in {"day", "week", "month"}:
                date_column = _resolve_column(all_headers, None, "date")
                keys = []
                for row in matched:
                    record_date = _as_date(row.values.get(date_column)) if date_column else None
                    if record_date is None:
                        keys.append("Unknown")
                    elif grouping == "day":
                        keys.append(record_date.isoformat())
                    elif grouping == "month":
                        keys.append(record_date.strftime("%Y-%m"))
                    else:
                        iso = record_date.isocalendar()
                        keys.append(f"{iso.year}-W{iso.week:02d}")
            else:
                group_column = _resolve_column(all_headers, grouping)
                keys = [str(row.values.get(group_column) or "Unknown") for row in matched] if group_column else ["All"] * len(matched)
            group_values: dict[str, list[float]] = {}
            for row, key in zip(matched, keys, strict=True):
                if metric == "count":
                    group_values.setdefault(key, []).append(1.0)
                else:
                    numeric = _as_number(row.values.get(target_column)) if target_column else None
                    if numeric is not None:
                        group_values.setdefault(key, []).append(numeric)
            grouped_results = {
                key: (sum(values) / len(values) if metric == "average" and values else sum(values))
                for key, values in group_values.items()
            }
        if operation == "aggregate":
            title = {"sum": "Tổng", "count": "Số lượng", "average": "Trung bình"}.get(metric, "Kết quả")
            answer = f"{title} {target_column or 'dữ liệu'}: {value:,.2f}".rstrip("0").rstrip(".")
        else:
            answer = f"Tìm thấy {len(matched)} dòng dữ liệu phù hợp trong {len({row.path for row in matched})} tệp."

    if not matched:
        answer = "Không tìm thấy dữ liệu phù hợp với yêu cầu trong các tệp đã chọn."
    return {
        "status": "column_not_found" if unresolved_target else "success" if matched else "no_results",
        "query": prompt,
        "answer": answer,
        "filters": {"salesperson": person, "start_date": start_date.isoformat() if start_date else None, "end_date": end_date.isoformat() if end_date else None},
        "operation": operation,
        "metric": metric if operation == "aggregate" else None,
        "target_column": target_column,
        "group_by": query_plan.get("group_by", "none"),
        "grouped_results": grouped_results if operation == "aggregate" else {},
        "total": total,
        "files_scanned": files_scanned,
        "rows_parsed": len(records),
        "matched_rows": len(matched),
        "available_columns": all_headers if unresolved_target else [],
        "data": {"columns": display_columns, "rows": display_rows},
        "warnings": warnings,
    }
