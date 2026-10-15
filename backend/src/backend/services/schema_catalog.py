"""Build and cache a header-only catalog for local Markdown and Excel files."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.services.local_vault import LocalVault
from backend.services.query_engine import _as_date, _clean_headers, _header_group, parse_markdown_records


_CACHE: dict[tuple[str, str, int, int], dict[str, Any]] = {}
_GROUP_ALIASES = {
    "date": ["date", "day", "ngày"],
    "person": ["person", "seller", "salesperson", "employee", "nhân viên"],
    "product": ["product", "item", "sản phẩm", "mặt hàng"],
    "amount": ["amount", "revenue", "sales", "doanh thu", "doanh số"],
}


def _header_index(rows: list[tuple[Any, ...]]) -> int | None:
    nonempty = [
        index for index, row in enumerate(rows)
        if any(value is not None and str(value).strip() for value in row)
    ]
    if not nonempty:
        return None
    first = nonempty[0]
    first_values = [value for value in rows[first] if value is not None and str(value).strip()]
    if len(first_values) >= 2 or any(_header_group(str(value)) for value in first_values):
        return first
    return next(
        (
            index for index in nonempty[1:]
            if any(_header_group(str(value)) for value in rows[index] if value is not None)
            or sum(value is not None and bool(str(value).strip()) for value in rows[index]) >= 2
        ),
        first,
    )


def _type_for(values: list[Any], group: str | None = None) -> str:
    values = [value for value in values if value is not None and str(value).strip()]
    if not values:
        return "unknown"
    if group == "date" and all(_as_date(value) is not None for value in values):
        return "date"
    if all(isinstance(value, bool) for value in values):
        return "boolean"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values):
        return "number"
    if all(isinstance(value, (date, datetime)) for value in values):
        return "date"
    return "text"


def _column_metadata(headers: list[str], rows: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    columns = []
    for column_index, name in enumerate(headers):
        values = [row[column_index] for row in rows if column_index < len(row)]
        group = _header_group(name)
        columns.append({
            "name": name,
            "type": _type_for(values[:200], group),
            "aliases": _GROUP_ALIASES.get(group, []),
        })
    return columns


def _catalog_xlsx(file_path: str, content: bytes) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [tuple(row) for row in worksheet.iter_rows(min_row=1, max_row=30, values_only=True)]
            index = _header_index(rows)
            if index is None:
                continue
            headers = _clean_headers(list(rows[index]))
            if not any(not value.startswith("Column ") for value in headers):
                continue
            sample_rows = [row for row in rows[index + 1:] if any(value is not None for value in row)]
            sheets.append({
                "name": worksheet.title,
                "columns": _column_metadata(headers, sample_rows),
                "row_count": max(0, worksheet.max_row - index - 1),
            })
    finally:
        workbook.close()
    return {"file_path": file_path, "sheets": sheets}


def _catalog_markdown(file_path: str, content: str) -> dict[str, Any]:
    records = parse_markdown_records(file_path, content)
    headers = list(dict.fromkeys(column for record in records for column in record.values))
    rows = [tuple(record.values.get(header) for header in headers) for record in records]
    sheets = []
    if headers:
        sheets.append({
            "name": "Markdown tables",
            "columns": _column_metadata(headers, rows),
            "row_count": len(records),
        })
    return {"file_path": file_path, "sheets": sheets}


def build_schema_catalog(vault: LocalVault) -> dict[str, Any]:
    """Catalog files below the selected root; cache unchanged file schemas."""
    files: list[dict[str, Any]] = []
    warnings: list[dict[str, str]] = []
    root_key = str(vault.root)
    active_keys: set[tuple[str, str, int, int]] = set()
    for relative in vault.list_files():
        path = vault.resolve_file(relative)
        stat = path.stat()
        key = (root_key, relative, stat.st_mtime_ns, stat.st_size)
        active_keys.add(key)
        cached = _CACHE.get(key)
        if cached is not None:
            files.append(cached)
            continue
        try:
            data = path.read_bytes()
            if path.suffix.lower() == ".xlsx":
                catalog_entry = _catalog_xlsx(relative, data)
            else:
                catalog_entry = _catalog_markdown(relative, data.decode("utf-8"))
            _CACHE[key] = catalog_entry
            files.append(catalog_entry)
        except Exception as error:
            warnings.append({"file": relative, "message": f"Schema could not be read ({type(error).__name__})."})
    for stale in [key for key in _CACHE if key[0] == root_key and key not in active_keys]:
        _CACHE.pop(stale, None)
    return {"files": files, "warnings": warnings}


def catalog_columns(catalog: dict[str, Any]) -> dict[tuple[str, str], list[dict[str, Any]]]:
    """Index exact column metadata by normalized source path and sheet name."""
    indexed: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for file_entry in catalog.get("files", []):
        for sheet in file_entry.get("sheets", []):
            indexed[(str(file_entry["file_path"]).casefold(), str(sheet["name"]).casefold())] = sheet["columns"]
    return indexed
