"""Load only planned sources and execute validated plans in DuckDB memory."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from backend.schemas.query import MultiFileQueryPlan, QueryStep
from backend.services.local_vault import LocalVault
from backend.services.query_compiler import InvalidQueryPlan, compile_step
from backend.services.query_engine import _clean_headers, parse_markdown_records
from backend.services.schema_catalog import _header_index, catalog_columns


MAX_FINAL_ROWS = 5000


class QueryExecutionError(RuntimeError):
    """A planned source could not be read or executed."""


class DuckDBUnavailableError(QueryExecutionError):
    """DuckDB is not installed in this backend environment."""


def _read_source(vault: LocalVault, step: QueryStep):
    import pandas as pd

    path = vault.resolve_file(step.source.file_path)
    content = path.read_bytes()
    if path.suffix.lower() == ".md":
        records = parse_markdown_records(step.source.file_path, content.decode("utf-8"))
        columns = list(dict.fromkeys(column for record in records for column in record.values))
        return pd.DataFrame([record.values for record in records], columns=columns)

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = next((sheet for sheet in workbook.worksheets if sheet.title == step.source.sheet_name), None)
        if worksheet is None:
            raise QueryExecutionError(f"Sheet '{step.source.sheet_name}' no longer exists")
        head = [tuple(row) for row in worksheet.iter_rows(min_row=1, max_row=30, values_only=True)]
        header_index = _header_index(head)
        if header_index is None:
            return pd.DataFrame()
        headers = _clean_headers(list(head[header_index]))
        rows = [
            tuple(row)
            for row in worksheet.iter_rows(min_row=header_index + 2, values_only=True)
            if any(value is not None and str(value).strip() for value in row)
        ]
        return pd.DataFrame.from_records(rows, columns=headers)
    finally:
        workbook.close()


def _coerce_source_types(frame: Any, columns: list[dict[str, Any]]) -> Any:
    """Convert spreadsheet strings using the types inferred by the catalog."""
    import pandas as pd

    for column in columns:
        name = str(column["name"])
        if name not in frame.columns:
            continue
        if column.get("type") == "date":
            frame[name] = pd.to_datetime(frame[name], errors="coerce", dayfirst=True)
        elif column.get("type") == "number":
            frame[name] = pd.to_numeric(frame[name], errors="coerce")
    return frame


def _json_value(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    try:
        import pandas as pd
        if pd.isna(value):
            return None
    except (ImportError, TypeError, ValueError):
        pass
    return value


def _data_type(series: Any) -> str:
    dtype = str(series.dtype).lower()
    if "bool" in dtype:
        return "boolean"
    if "int" in dtype or "float" in dtype or "decimal" in dtype:
        return "number"
    if "date" in dtype or "time" in dtype:
        return "date"
    sample = next((value for value in series.tolist() if value is not None), None)
    if isinstance(sample, (datetime, date)):
        return "date"
    if isinstance(sample, bool):
        return "boolean"
    if isinstance(sample, (int, float)):
        return "number"
    return "text"


def _describe_result(frame: Any) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    columns = [{"name": str(name), "type": _data_type(frame[name])} for name in frame.columns]
    rows = [
        {str(column): _json_value(value) for column, value in record.items()}
        for record in frame.to_dict(orient="records")
    ]
    return columns, rows


def _answer(plan: MultiFileQueryPlan, columns: list[dict[str, str]], rows: list[dict[str, Any]]) -> str:
    if not rows:
        return "Không tìm thấy dòng dữ liệu phù hợp trong các tệp đã chọn."
    if len(rows) == 1 and len(columns) == 1:
        column = columns[0]["name"]
        return f"{column}: {rows[0].get(column)}"
    return f"Tìm thấy {len(rows)} dòng dữ liệu phù hợp từ {len(plan.steps)} bước truy vấn."


def execute_plan(
    vault: LocalVault,
    plan: MultiFileQueryPlan,
    catalog: dict[str, Any],
) -> dict[str, Any]:
    """Execute the multi-step plan locally and return generic columns and rows."""
    try:
        import duckdb
    except ImportError as error:
        raise DuckDBUnavailableError("DuckDB is missing; run `uv sync` from backend/ to install dependencies") from error

    column_index = catalog_columns(catalog)
    connection = duckdb.connect(database=":memory:")
    warnings: list[dict[str, str]] = list(catalog.get("warnings", []))
    step_frames: dict[str, Any] = {}
    try:
        for index, step in enumerate(plan.steps):
            source_key = (step.source.file_path.replace("\\", "/").casefold(), step.source.sheet_name.casefold())
            try:
                source_columns = column_index[source_key]
                frame = _coerce_source_types(_read_source(vault, step), source_columns)
                relation = f"source_{index}"
                connection.register(relation, frame)
                sql, parameters = compile_step(step, relation, source_columns)
                frame_result = connection.execute(sql, parameters).fetchdf()
            except Exception as error:
                if isinstance(error, (InvalidQueryPlan, QueryExecutionError)):
                    raise
                raise QueryExecutionError(
                    f"Could not execute step '{step.id}' for {step.source.file_path}#{step.source.sheet_name}"
                ) from error
            step_frames[step.id] = frame_result

        if plan.combine.operation == "none":
            final = step_frames[plan.steps[0].id]
        else:
            combine = plan.combine
            left_name = "step_left"
            right_name = "step_right"
            connection.register(left_name, step_frames[combine.left_step or ""])
            connection.register(right_name, step_frames[combine.right_step or ""])
            if combine.operation == "union":
                combined_sql = (
                    f'SELECT * FROM "{left_name}" UNION ALL SELECT * FROM "{right_name}" '
                    f"LIMIT {MAX_FINAL_ROWS + 1}"
                )
            else:
                left_columns = list(step_frames[combine.left_step or ""].columns)
                right_columns = list(step_frames[combine.right_step or ""].columns)
                right_extra = [column for column in right_columns if column != combine.right_column]
                duplicates = set(left_columns).intersection(right_extra)
                if duplicates:
                    raise InvalidQueryPlan(
                        "Join results contain duplicate column names: " + ", ".join(sorted(duplicates))
                    )
                left_key = '"' + str(combine.left_column).replace('"', '""') + '"'
                right_key = '"' + str(combine.right_column).replace('"', '""') + '"'
                right_projection = ", ".join(
                    f'r."{column.replace(chr(34), chr(34) * 2)}"' for column in right_extra
                )
                projection = "l.*" + (f", {right_projection}" if right_projection else "")
                combined_sql = (
                    f'SELECT {projection} FROM "{left_name}" l '
                    f'{combine.join_type.upper()} JOIN "{right_name}" r '
                    f"ON l.{left_key} = r.{right_key} LIMIT {MAX_FINAL_ROWS + 1}"
                )
            final = connection.execute(combined_sql).fetchdf()

        if plan.display.columns:
            missing = [name for name in plan.display.columns if name not in final.columns]
            if missing:
                raise InvalidQueryPlan(f"Display columns are unavailable: {', '.join(missing)}")
            final = final.loc[:, plan.display.columns]
        truncated = len(final) > MAX_FINAL_ROWS
        if truncated:
            final = final.iloc[:MAX_FINAL_ROWS]
            warnings.append({"file": "query", "message": f"Result was limited to {MAX_FINAL_ROWS} rows."})
        result_columns, result_rows = _describe_result(final)
        sources = list(dict.fromkeys(
            f"{step.source.file_path}#{step.source.sheet_name}" for step in plan.steps
        ))
        return {
            "status": "success" if result_rows else "no_results",
            "query": "",
            "answer": _answer(plan, result_columns, result_rows),
            "plan": plan.model_dump(mode="json", by_alias=True),
            "display": plan.display.model_dump(mode="json"),
            "result": {
                "columns": result_columns,
                "rows": result_rows,
                "row_count": len(result_rows),
                "truncated": truncated,
            },
            "step_results": [
                {"step_id": step.id, "row_count": len(step_frames[step.id])}
                for step in plan.steps
            ],
            "sources": sources,
            "warnings": warnings,
        }
    finally:
        connection.close()
