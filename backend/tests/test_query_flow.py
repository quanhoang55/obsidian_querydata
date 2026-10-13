from __future__ import annotations

import json
import unittest
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx
from google.genai.errors import APIError
from openpyxl import Workbook

from backend.clients.gemini import GeminiClient, GeminiInvalidResponseError, GeminiTimeoutError
from backend.clients.gemini import GeminiClientError
from backend.config import Settings
from backend.schemas.query import QueryPlan, gemini_response_schema
from backend.services import query_service
from backend.services.query_engine import aggregate_sales, parse_excel_records


def make_plan(**overrides: object) -> QueryPlan:
    values: dict[str, object] = {
        "operation": "aggregate",
        "metric": "sum",
        "target_column": "Doanh thu",
        "display_columns": ["Ngày bán", "Người bán", "Doanh thu"],
        "person": "Lan",
        "start_date": date(2026, 10, 5),
        "end_date": date(2026, 10, 11),
        "group_by": "none",
        "file_path": None,
        "clarification_needed": False,
        "clarification_question": None,
    }
    values.update(overrides)
    return QueryPlan.model_validate(values)


def make_settings(**overrides: object) -> Settings:
    values: dict[str, object] = {
        "obsidian_url": "https://127.0.0.1:27124",
        "obsidian_key": "obsidian-test-key",
        "gemini_api_key": "gemini-test-key",
    }
    values.update(overrides)
    return Settings(**values)


class FakeGeminiClient:
    def __init__(self, plan: QueryPlan) -> None:
        self.plan = plan
        self.close = AsyncMock()

    async def create_query_plan(self, prompt: str) -> QueryPlan:
        self.prompt = prompt
        return self.plan


class QueryPlanTests(unittest.TestCase):
    def test_rejects_unknown_fields_and_invalid_dates(self) -> None:
        with self.assertRaises(ValueError):
            QueryPlan.model_validate({**make_plan().model_dump(), "unexpected": "value"})
        with self.assertRaises(ValueError):
            make_plan(start_date="2026-02-30")

    def test_rejects_reversed_date_range(self) -> None:
        with self.assertRaises(ValueError):
            make_plan(start_date=date(2026, 10, 12), end_date=date(2026, 10, 11))


class SpreadsheetQueryTests(unittest.TestCase):
    @staticmethod
    def make_workbook() -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Ngày bán", "Mặt hàng", "Người bán", "Doanh thu", "Khu vực"])
        sheet.append([date(2026, 10, 11), "Cà phê", "Lan", 120000, "Hà Nội"])
        sheet.append([date(2026, 10, 11), "Trà", "Minh", 80000, "Đà Nẵng"])
        sheet.append([date(2026, 10, 10), "Bánh", "Lan", 50000, "Hà Nội"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    def test_parses_xlsx_and_aggregates_requested_day(self) -> None:
        content = self.make_workbook()
        records = parse_excel_records("Test_xlsx/test_query.xlsx", content)
        self.assertEqual(len(records), 3)
        result = aggregate_sales(
            "trong ngày 11/10/2026, tổng doanh thu của tất cả người bán; tìm mặt hàng trong file Test_xlsx/test_query.xlsx",
            [("Test_xlsx/test_query.xlsx", content), ("Other/data.xlsx", content)],
            make_plan(person=None, start_date=date(2026, 10, 11), end_date=date(2026, 10, 11), display_columns=["Ngày bán", "Mặt hàng", "Người bán", "Doanh thu", "Khu vực"], file_path="Test_xlsx/test_query.xlsx").model_dump(mode="json"),
        )
        self.assertEqual(result["total"], 200000)
        self.assertEqual(result["matched_rows"], 2)
        self.assertEqual(result["data"]["columns"], ["Ngày bán", "Mặt hàng", "Người bán", "Doanh thu", "Khu vực"])
        self.assertEqual(result["data"]["rows"][0]["values"], {
            "Ngày bán": datetime(2026, 10, 11),
            "Mặt hàng": "Cà phê",
            "Người bán": "Lan",
            "Doanh thu": 120000,
            "Khu vực": "Hà Nội",
        })
        self.assertEqual(result["files_scanned"], 1)
        self.assertEqual(result["warnings"], [])

    def test_distinct_intent_lists_dynamic_column_values(self) -> None:
        content = self.make_workbook()
        result = aggregate_sales(
            "Tất cả các nhân viên sale trong file Test_xlsx/test_query.xlsx",
            [("Test_xlsx/test_query.xlsx", content)],
            make_plan(
                operation="distinct", metric=None, target_column="salesperson",
                display_columns=["salesperson"], person=None,
                start_date=None, end_date=None, file_path="Test_xlsx/test_query.xlsx",
            ).model_dump(mode="json"),
        )
        self.assertEqual(result["target_column"], "Người bán")
        self.assertEqual([row["values"]["Người bán"] for row in result["data"]["rows"]], ["Lan", "Minh"])

    def test_distinct_intent_supports_a_non_sales_workbook_schema(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Department", "Priority"])
        sheet.append(["Support", "High"])
        sheet.append(["Sales", "Medium"])
        sheet.append(["Support", "Low"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        result = aggregate_sales(
            "List all departments in the workbook",
            [("Operations/teams.xlsx", output.getvalue())],
            make_plan(
                operation="distinct", metric=None, target_column="Department",
                display_columns=["Department"], person=None, start_date=None,
                end_date=None, file_path="Operations/teams.xlsx",
            ).model_dump(mode="json"),
        )
        self.assertEqual(result["data"]["columns"], ["Department"])
        self.assertEqual([row["values"]["Department"] for row in result["data"]["rows"]], ["Support", "Sales"])

    def test_resolves_seller_intent_to_common_abbreviated_header(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Tên NVKD", "Revenue"])
        sheet.append(["Lan", 120000])
        sheet.append(["Minh", 80000])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        result = aggregate_sales(
            "List all sales people",
            [("Sales/data.xlsx", output.getvalue())],
            make_plan(
                operation="distinct", metric=None, target_column="seller",
                display_columns=["seller"], person=None, start_date=None,
                end_date=None, file_path="Sales/data.xlsx",
            ).model_dump(mode="json"),
        )
        self.assertEqual(result["target_column"], "Tên NVKD")
        self.assertEqual([row["values"]["Tên NVKD"] for row in result["data"]["rows"]], ["Lan", "Minh"])


class QueryServiceTests(unittest.IsolatedAsyncioTestCase):
    async def test_invalid_gemini_response_does_not_load_vault(self) -> None:
        fake_gemini = FakeGeminiClient(make_plan())
        fake_gemini.create_query_plan = AsyncMock(side_effect=GeminiInvalidResponseError("bad plan"))
        with (
            patch.object(query_service, "GeminiClient", return_value=fake_gemini),
            patch.object(query_service, "_load_vault_documents", new_callable=AsyncMock) as load_docs,
        ):
            with self.assertRaises(GeminiInvalidResponseError):
                await query_service.query_sales("query", make_settings())

        load_docs.assert_not_awaited()
        fake_gemini.close.assert_awaited_once()

    async def test_clarification_does_not_load_vault(self) -> None:
        plan = make_plan(
            person=None,
            start_date=None,
            end_date=None,
            clarification_needed=True,
            clarification_question="Which date range should I use?",
        )
        fake_gemini = FakeGeminiClient(plan)
        with (
            patch.object(query_service, "GeminiClient", return_value=fake_gemini),
            patch.object(query_service, "_load_vault_documents", new_callable=AsyncMock) as load_docs,
        ):
            result = await query_service.query_sales("Compare sales", make_settings())

        load_docs.assert_not_awaited()
        self.assertEqual(result["status"], "clarification_required")
        self.assertEqual(result["files_scanned"], 0)
        self.assertEqual(result["answer"], "Which date range should I use?")
        fake_gemini.close.assert_awaited_once()

    async def test_executes_typed_plan_against_locally_loaded_rows(self) -> None:
        fake_gemini = FakeGeminiClient(make_plan())
        markdown = """| Ngày bán | Nhân viên | Doanh thu |
| --- | --- | --- |
| 2026-10-07 | Lan | 120000 |
| 2026-10-08 | Lan | 80000 |
| 2026-10-09 | Minh | 90000 |
"""
        with (
            patch.object(query_service, "GeminiClient", return_value=fake_gemini),
            patch.object(
                query_service,
                "_load_vault_documents",
                new=AsyncMock(return_value=[("Sales/October.md", markdown)]),
            ),
        ):
            result = await query_service.query_sales("Long prompt asking about Lan", make_settings())

        self.assertEqual(result["total"], 200000)
        self.assertEqual(result["matched_rows"], 2)
        self.assertEqual(result["query_plan"]["person"], "Lan")
        self.assertEqual(result["filters"]["start_date"], "2026-10-05")


class GeminiClientTests(unittest.IsolatedAsyncioTestCase):
    async def test_parses_structured_json_and_sends_prompt_only(self) -> None:
        response_body = {
            "operation": "aggregate",
            "metric": "sum",
            "target_column": "Doanh thu",
            "display_columns": ["Người bán", "Doanh thu"],
            "person": "Lan",
            "start_date": "2026-10-05",
            "end_date": "2026-10-11",
            "group_by": "none",
            "file_path": None,
            "clarification_needed": False,
            "clarification_question": None,
        }
        create_interaction = AsyncMock(return_value=SimpleNamespace(output_text=json.dumps(response_body)))
        client = GeminiClient(make_settings())
        client._client = SimpleNamespace(
            aio=SimpleNamespace(
                interactions=SimpleNamespace(create=create_interaction),
                aclose=AsyncMock(),
            )
        )

        plan = await client.create_query_plan("Doanh số của Lan tuần trước")

        self.assertEqual(plan.person, "Lan")
        request = create_interaction.await_args.kwargs
        self.assertIn("Doanh số của Lan tuần trước", request["input"])
        self.assertNotIn("vault content", request["input"].casefold())
        request_schema = request["response_format"]["schema"]
        self.assertEqual(request_schema["type"], "OBJECT")
        self.assertNotIn("additionalProperties", request_schema)
        self.assertEqual(request_schema["properties"]["operation"]["enum"], ["aggregate", "distinct", "inspect"])
        self.assertTrue(request_schema["properties"]["person"]["nullable"])
        self.assertEqual(gemini_response_schema(), request_schema)
        self.assertEqual(request["response_format"]["mime_type"], "application/json")
        self.assertEqual(request["response_format"]["type"], "text")
        await client.close()

    async def test_rejects_malformed_json_plan(self) -> None:
        create_interaction = AsyncMock(return_value=SimpleNamespace(output_text='{"metric":"unknown"}'))
        client = GeminiClient(make_settings())
        client._client = SimpleNamespace(
            aio=SimpleNamespace(
                interactions=SimpleNamespace(create=create_interaction),
                aclose=AsyncMock(),
            )
        )

        with self.assertRaises(GeminiInvalidResponseError):
            await client.create_query_plan("show sales")
        await client.close()

    async def test_reports_network_failure_clearly(self) -> None:
        create_interaction = AsyncMock(side_effect=httpx.ConnectError("connection refused"))
        client = GeminiClient(make_settings())
        client._client = SimpleNamespace(
            aio=SimpleNamespace(
                interactions=SimpleNamespace(create=create_interaction),
                aclose=AsyncMock(),
            )
        )

        with self.assertRaisesRegex(GeminiClientError, "Could not connect to the Gemini API"):
            await client.create_query_plan("show sales")
        await client.close()

    async def test_reports_provider_rejection_and_redacts_key(self) -> None:
        api_key = "test-secret-key"
        provider_error = APIError(400, {"error": {"message": f"Invalid request for {api_key}", "status": "INVALID_ARGUMENT"}})
        create_interaction = AsyncMock(side_effect=provider_error)
        client = GeminiClient(make_settings(gemini_api_key=api_key))
        client._client = SimpleNamespace(
            aio=SimpleNamespace(
                interactions=SimpleNamespace(create=create_interaction),
                aclose=AsyncMock(),
            )
        )

        with self.assertRaises(GeminiClientError) as raised:
            await client.create_query_plan("show sales")
        self.assertIn("HTTP 400 (INVALID_ARGUMENT)", str(raised.exception))
        self.assertNotIn(api_key, str(raised.exception))
        await client.close()


class QueryEndpointTests(unittest.IsolatedAsyncioTestCase):
    async def test_gemini_timeout_maps_to_504(self) -> None:
        from fastapi import HTTPException

        from backend.api.endpoints import routes

        with (
            patch.object(routes, "get_settings", return_value=make_settings()),
            patch.object(routes, "query_sales", new=AsyncMock(side_effect=GeminiTimeoutError("timeout"))),
        ):
            with self.assertRaises(HTTPException) as raised:
                await routes.query_vault(routes.QueryRequest(prompt="sales"))

        self.assertEqual(raised.exception.status_code, 504)

    async def test_missing_gemini_key_maps_to_503_before_service_call(self) -> None:
        from fastapi import HTTPException

        from backend.api.endpoints import routes

        with (
            patch.object(routes, "get_settings", return_value=make_settings(gemini_api_key="")),
            patch.object(routes, "query_sales", new=AsyncMock()) as query,
        ):
            with self.assertRaises(HTTPException) as raised:
                await routes.query_vault(routes.QueryRequest(prompt="sales"))

        self.assertEqual(raised.exception.status_code, 503)
        query.assert_not_awaited()


if __name__ == "__main__":
    unittest.main()
